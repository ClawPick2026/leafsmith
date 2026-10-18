"""Real generated fixtures; pypdf/PDFium/Office independent read-back."""
import importlib.util
import json
from pathlib import Path
import pytest
from reportlab.pdfgen import canvas
from pypdf import PdfReader

MODULE = Path(__file__).parents[1] / 'pdf_ops.py'

def run(tmp_path, operation, inputs=(), engines=None, **options):
    if not MODULE.exists():
        pytest.fail('native PDF worker not implemented')
    spec = importlib.util.spec_from_file_location('pdf_ops', MODULE)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.run({'operation':operation,'inputs':[str(p) for p in inputs],
                       'output_dir':str(tmp_path / 'out'),'options':options,'engines':engines or {}})

@pytest.fixture
def source(tmp_path):
    p = tmp_path / 'source.pdf'
    c = canvas.Canvas(str(p), pageCompression=0)
    c.drawString(72, 740, 'Original text')
    c.drawString(72, 700, 'DRAFT')
    c.save()
    return p

def output(result):
    json.dumps(result)
    assert len(result['outputs']) >= 1
    p = Path(result['outputs'][0]['path'])
    assert p.is_file()
    assert p.name == result['outputs'][0]['name']
    return p

def test_edit_replaces_object_not_overlay(tmp_path, source):
    original = source.read_bytes()
    info = run(tmp_path, 'inspect_pdf', [source])
    index = next(o['object_index'] for o in info['details']['pages'][0]['objects'] if o.get('text')=='Original text')
    p = output(run(tmp_path, 'edit_pdf_text', [source], page=1, object_index=index,
                   text='Replacement', font='Courier', font_size=18))
    text = PdfReader(p).pages[0].extract_text()
    assert 'Replacement' in text and 'Original text' not in text and 'DRAFT' in text
    after = run(tmp_path, 'inspect_pdf', [p])['details']['pages'][0]['objects']
    assert len(after) == len(info['details']['pages'][0]['objects'])
    assert next(o for o in after if o.get('text')=='Replacement')['font_size'] == 18
    assert source.read_bytes() == original


def test_password_security_roundtrip(tmp_path, source):
    p = output(run(tmp_path, 'protect_pdf', [source], password='user-secret', owner_password='owner-secret'))
    reader = PdfReader(p)
    assert reader.is_encrypted and reader.trailer['/Encrypt']['/V'] == 5
    assert reader.trailer['/Encrypt']['/Length'] == 256
    assert not reader.decrypt('wrong')
    with pytest.raises(ValueError, match='password'):
        run(tmp_path, 'unlock_pdf', [p], password='wrong')
    clear = output(run(tmp_path, 'unlock_pdf', [p], password='user-secret'))
    assert not PdfReader(clear).is_encrypted
    assert 'Original text' in PdfReader(clear).pages[0].extract_text()


def test_remove_selected_exact_watermark(tmp_path, source):
    p = output(run(tmp_path, 'remove_watermark', [source], exact_text='DRAFT', page=1))
    text = PdfReader(p).pages[0].extract_text()
    assert 'DRAFT' not in text and 'Original text' in text
    assert len(run(tmp_path,'inspect_pdf',[p])['details']['pages'][0]['objects']) == 1
    with pytest.raises(ValueError, match='match'):
        run(tmp_path, 'remove_watermark', [source], exact_text='DRA')


@pytest.mark.parametrize('template_id,fields', [('contact',{'name','email','message'}),('approval',{'requester','description','approved'}),('invoice',{'invoice_number','customer','amount'})])
def test_form_template_has_real_fields(tmp_path, template_id, fields):
    p = output(run(tmp_path,'form_template',template_id=template_id))
    reader = PdfReader(p)
    assert fields <= set(reader.get_fields())
    assert all(v['/FT'] in {'/Tx','/Btn'} for v in reader.get_fields().values())
    import pypdfium2
    with pypdfium2.PdfDocument(str(p)) as doc:
        page=doc[0]; bitmap=page.render(scale=1); image=bitmap.to_pil()
        assert image.width > 500
        bitmap.close(); page.close()


@pytest.mark.parametrize('kind,subtype',[('comment','/Text'),('highlight','/Highlight')])
def test_native_annotations(tmp_path, source, kind, subtype):
    p = output(run(tmp_path,'annotate_pdf',[source],page=1,kind=kind,text='Review this',rect=[72,730,180,755]))
    annotation = PdfReader(p).pages[0]['/Annots'][0].get_object()
    assert annotation['/Subtype'] == subtype
    assert annotation['/Contents'] == 'Review this'
    assert 'Original text' in PdfReader(p).pages[0].extract_text()
    if kind == 'highlight': assert len(annotation['/QuadPoints']) == 8


@pytest.mark.parametrize('mode',['lossless','raster'])
def test_compression_real_readback(tmp_path, source, mode):
    p = output(run(tmp_path,'compress_pdf',[source],mode=mode,dpi=100,quality=60))
    reader=PdfReader(p)
    assert len(reader.pages)==1
    if mode=='lossless':
        assert 'Original text' in reader.pages[0].extract_text()
        assert p.stat().st_size < source.stat().st_size
    else:
        assert not reader.pages[0].extract_text().strip()
        assert reader.pages[0].images


@pytest.mark.parametrize('format',['txt','html','csv','docx','xlsx','pptx','png','jpeg'])
def test_export_real_artifacts(tmp_path, source, format):
    p = output(run(tmp_path,'convert_from_pdf',[source],format=format,dpi=100))
    if format in {'txt','html','csv'}: assert 'Original text' in p.read_text()
    elif format=='docx':
        from docx import Document
        assert 'Original text' in '\n'.join(v.text for v in Document(p).paragraphs)
    elif format=='xlsx':
        from openpyxl import load_workbook
        assert any('Original text' in str(cell.value) for sheet in load_workbook(p) for row in sheet for cell in row)
    elif format=='pptx':
        from pptx import Presentation
        assert any('Original text' in s.text for slide in Presentation(p).slides for s in slide.shapes if s.has_text_frame)
    else:
        from PIL import Image
        with Image.open(p) as image: assert image.width > 500 and image.height > 500


@pytest.mark.parametrize('kind',['txt','html','png'])
def test_local_to_pdf(tmp_path, kind):
    p=tmp_path / ('input.'+kind)
    if kind=='png':
        from PIL import Image
        Image.new('RGB',(320,180),'navy').save(p)
    elif kind=='html':
        p.write_text('<h1>Safe heading</h1><script>secret_active_script</script><img src="http://127.0.0.1:9/no"><p>Local text</p>')
    else: p.write_text('Local text\nSecond line')
    result=run(tmp_path,'convert_to_pdf',[p])
    pdf=PdfReader(output(result))
    assert len(pdf.pages)==1
    if kind=='png': assert pdf.pages[0].images
    else:
        assert 'Local text' in pdf.pages[0].extract_text()
        assert 'secret_active_script' not in pdf.pages[0].extract_text()


def test_office_missing_engine_is_actionable(tmp_path):
    from docx import Document
    p=tmp_path/'sample.docx'; doc=Document(); doc.add_paragraph('Office fixture'); doc.save(p)
    with pytest.raises(ValueError, match='Install the optional LibreOffice component'):
        run(tmp_path,'convert_to_pdf',[p])


def test_unsafe_office_archive_rejected(tmp_path):
    import zipfile
    p=tmp_path/'unsafe.docx'
    with zipfile.ZipFile(p,'w') as archive: archive.writestr('../escape','blocked')
    with pytest.raises(ValueError, match='Unsafe Office archive'):
        run(tmp_path,'convert_to_pdf',[p])


@pytest.mark.parametrize('operation',['ocr_pdf','ocr_image'])
def test_ocr_missing_engine_actionable(tmp_path, source, operation):
    if operation=='ocr_image':
        from PIL import Image
        source=tmp_path/'scan.png'; Image.new('RGB',(400,120),'white').save(source)
    with pytest.raises(ValueError, match='Install the optional Tesseract component'):
        run(tmp_path,operation,[source])


def test_ocr_failed_executable_never_returns_success(tmp_path, source):
    import sys
    with pytest.raises(ValueError, match='Tesseract failed'):
        run(tmp_path,'ocr_pdf',[source],engines={'tesseract':sys.executable})


@pytest.mark.parametrize('operation',['ocr_pdf','ocr_image'])
def test_optional_real_ocr(tmp_path, operation):
    import os
    engine=os.environ.get('PAPERWORK_TESSERACT')
    if not engine: pytest.skip('Tesseract absent; real OCR not verified on this host')
    from PIL import Image, ImageDraw, ImageFont
    image=Image.new('RGB',(1400,320),'white')
    draw=ImageDraw.Draw(image)
    font=ImageFont.load_default(size=70)
    draw.text((60,90),'PAPERWORK OCR 12345',font=font,fill='black')
    p=tmp_path/'scan.png'; image.save(p)
    if operation=='ocr_pdf':
        p=tmp_path/'scan.pdf'; image.save(p,format='PDF',resolution=150)
    r=run(tmp_path,operation,[p],engines={'tesseract':engine},dpi=200)
    assert len(r['outputs'])==2
    pdf=next(Path(o['path']) for o in r['outputs'] if o['mime']=='application/pdf')
    txt=next(Path(o['path']) for o in r['outputs'] if o['mime']=='text/plain')
    assert 'PAPERWORK' in PdfReader(pdf).pages[0].extract_text().upper()
    assert 'PAPERWORK' in txt.read_text().upper()


def test_invalid_options_and_paths_rejected_without_outputs(tmp_path, source):
    with pytest.raises(ValueError,match='Unknown option'):
        run(tmp_path,'compress_pdf',[source],misspelled_mode='raster')
    with pytest.raises(ValueError,match='page'):
        run(tmp_path,'edit_pdf_text',[source],page=0,object_index=0,text='x')
    with pytest.raises(ValueError,match='page'):
        run(tmp_path,'inspect_pdf',[source],page=0)
    out=tmp_path/'out'
    assert not list(out.glob('paperwork-*'))


def test_missing_engine_leaves_no_output_files(tmp_path):
    from docx import Document
    p=tmp_path/'sample.docx'; Document().save(p)
    with pytest.raises(ValueError): run(tmp_path,'convert_to_pdf',[p])
    assert not list((tmp_path/'out').glob('paperwork-*'))


def test_office_engine_receives_locked_down_job_profile(tmp_path):
    import sys
    import xml.etree.ElementTree as ET
    from docx import Document
    p=tmp_path/'sample.docx'; Document().save(p)
    capture=tmp_path/'profile-observed.xml'
    probe=tmp_path/'office-probe'
    probe.write_text('#!'+sys.executable+'\nimport sys\nfrom pathlib import Path\nfrom urllib.parse import urlparse,unquote\nuri=next(a.split("=",1)[1] for a in sys.argv if a.startswith("-env:UserInstallation="))\nprofile=Path(unquote(urlparse(uri).path))\nPath('+repr(str(capture))+').write_bytes((profile/"user"/"registrymodifications.xcu").read_bytes())\nsys.exit(1)\n')
    probe.chmod(0o700)
    with pytest.raises(ValueError,match='LibreOffice failed'):
        run(tmp_path,'convert_to_pdf',[p],engines={'libreoffice':str(probe)})
    root=ET.fromstring(capture.read_text())
    props={n.attrib['{http://openoffice.org/2001/registry}name']:n.find('value').text for n in root.iter('prop')}
    assert props['MacroSecurityLevel']=='3'
    assert props['DisableMacrosExecution']=='true'
    # Headless conversion uses UNO UpdateDocMode=ACCORDING_TO_CONFIG. Actual schema enums:
    # Writer Content/Update/Link=2 Never; Calc Content/Update/Link=1 Never.
    # UpdateDocMode is NOT a Common/Load registry property.
    paths={n.attrib['{http://openoffice.org/2001/registry}path']:n for n in root.iter('item')}
    assert paths['/org.openoffice.Office.Writer/Content/Update'].find('prop/value').text=='2'
    assert paths['/org.openoffice.Office.Calc/Content/Update'].find('prop/value').text=='1'
    assert not list((tmp_path/'out').iterdir())


def test_edit_preserves_font_and_rejects_clipping_text(tmp_path, source):
    p=output(run(tmp_path,'edit_pdf_text',[source],page=1,object_index=0,text='Changed',font_size=14))
    assert 'Changed' in PdfReader(p).pages[0].extract_text()
    clipped=tmp_path/'clip.pdf'
    c=canvas.Canvas(str(clipped)); t=c.beginText(72,700); t.setTextRenderMode(7); t.textOut('Clip text'); c.drawText(t); c.save()
    with pytest.raises(ValueError,match='complex'):
        run(tmp_path,'edit_pdf_text',[clipped],object_index=0,text='Changed',font='Helvetica')
    info=run(tmp_path,'inspect_pdf',[clipped])['details']['pages'][0]['objects'][0]
    assert info['editable'] is False
    assert info['font']


@pytest.mark.parametrize('format',['docx','xlsx','pptx','odt','ods','odp'])
def test_optional_real_office_import(tmp_path, format):
    import os, zipfile
    engine=os.environ.get('PAPERWORK_LIBREOFFICE')
    if not engine: pytest.skip('LibreOffice absent; real Office import not verified on this host')
    p=tmp_path/('office.'+format)
    if format=='docx':
        from docx import Document
        doc=Document(); doc.add_paragraph('Office fixture'); doc.save(p)
    elif format=='xlsx':
        from openpyxl import Workbook
        doc=Workbook(); doc.active['A1']='Office fixture'; doc.save(p)
    elif format=='pptx':
        from pptx import Presentation
        from pptx.util import Inches
        doc=Presentation(); slide=doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='Office fixture'; doc.save(p)
    else:
        containers={'odt':('text','<text:p>Office fixture</text:p>','text'),
                    'ods':('spreadsheet','<table:table table:name="Sheet1"><table:table-row><table:table-cell office:value-type="string"><text:p>Office fixture</text:p></table:table-cell></table:table-row></table:table>','spreadsheet'),
                    'odp':('presentation','<draw:page draw:name="slide1"><draw:frame svg:x="2cm" svg:y="2cm" svg:width="10cm" svg:height="2cm"><draw:text-box><text:p>Office fixture</text:p></draw:text-box></draw:frame></draw:page>','presentation')}
        tag,body,mime=containers[format]
        document='<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.2"><office:body><office:'+tag+'>'+body+'</office:'+tag+'></office:body></office:document-content>'
        with zipfile.ZipFile(p,'w') as z:
            z.writestr('mimetype','application/vnd.oasis.opendocument.'+mime)
            z.writestr('content.xml',document)
            z.writestr('META-INF/manifest.xml','<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.'+mime+'"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    pdf=output(run(tmp_path,'convert_to_pdf',[p],engines={'libreoffice':engine}))
    assert 'Office fixture' in '\n'.join(page.extract_text() for page in PdfReader(pdf).pages)


def test_selected_annotation_removal_preserves_text(tmp_path, source):
    p=output(run(tmp_path,'annotate_pdf',[source],kind='comment',text='Selected',rect=[20,20,60,60]))
    r=run(tmp_path,'remove_watermark',[p],annotation_subtype='Text')
    clear=PdfReader(output(r))
    assert not clear.pages[0].get('/Annots')
    assert 'Original text' in clear.pages[0].extract_text()
    assert r['details']['removed_objects']==1


def test_table_reconstruction_csv_and_xlsx(tmp_path):
    from reportlab.platypus import SimpleDocTemplate, Table
    from reportlab.lib import colors
    import csv
    from openpyxl import load_workbook
    p=tmp_path/'table.pdf'
    table=Table([['Item','Amount'],['Widget','42'],['=UNTRUSTED','7']],colWidths=[180,100])
    table.setStyle([('GRID',(0,0),(-1,-1),1,colors.black)])
    SimpleDocTemplate(str(p)).build([table])
    exported=output(run(tmp_path,'convert_from_pdf',[p],format='csv'))
    with exported.open() as file:
        rows=list(csv.reader(file))
    assert rows[1]==['Widget','42'] and rows[2][0]=="'=UNTRUSTED"
    book=load_workbook(output(run(tmp_path,'convert_from_pdf',[p],format='xlsx')))
    assert book.active['A2'].value=='Widget' and book.active['B2'].value=='42'
    assert book.active['A3'].data_type!='f'


def test_subprocess_timeout_sanitized(tmp_path, source):
    import sys
    probe=tmp_path/'sleep-probe'
    probe.write_text('#!'+sys.executable+'\nimport time\ntime.sleep(5)\n')
    probe.chmod(0o700)
    with pytest.raises(ValueError,match='Tesseract timed out'):
        run(tmp_path,'ocr_pdf',[source],timeout=1,engines={'tesseract':str(probe)})
    assert not list((tmp_path/'out').iterdir())


def test_complex_script_edit_errors_instead_of_claiming_success(tmp_path, source):
    with pytest.raises(ValueError,match='Complex-script'):
        run(tmp_path,'edit_pdf_text',[source],object_index=0,text='مرحبا')


def test_inspect_real_objects(tmp_path, source):
    r = run(tmp_path, 'inspect_pdf', [source])
    pages = r['details']['pages']
    assert len(pages) == 1
    assert pages[0]['page'] == 1
    assert any(o['text'] == 'Original text' for o in pages[0]['objects'] if o['type']=='text')
    assert r['outputs'] == []
