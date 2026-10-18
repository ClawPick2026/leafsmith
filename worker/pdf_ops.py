"""Local native PDF operations. No network, shell, or source-file overwrites."""
from pathlib import Path
from contextlib import closing
import functools
import mimetypes
import os
import tempfile
import pypdfium2 as pdfium
import pypdfium2.raw as raw

OPERATIONS = {'inspect_pdf','edit_pdf_text','convert_to_pdf','convert_from_pdf',
              'form_template','protect_pdf','unlock_pdf','remove_watermark',
              'compress_pdf','ocr_pdf','ocr_image','annotate_pdf'}

class PdfOperationError(ValueError):
    """Safe, actionable errors suitable for the JSON protocol."""

def _inputs(request):
    values = request.get('inputs', [])
    if not isinstance(values, list):
        raise PdfOperationError('inputs must be a list of absolute file paths.')
    paths = []
    for value in values:
        p = Path(value)
        if not p.is_absolute() or not p.is_file() or p.is_symlink():
            raise PdfOperationError('Input must be an existing absolute regular file, not a symlink.')
        if p.stat().st_size > 512 * 1024 * 1024:
            raise PdfOperationError('Input exceeds the 512 MiB worker limit.')
        paths.append(p.resolve())
    return paths

def _objects(page):
    with closing(page.get_textpage()) as textpage:
        result = []
        for index, obj in enumerate(page.get_objects(max_depth=1, textpage=textpage)):
            item = {'object_index':index, 'type':'text' if obj.type == raw.FPDF_PAGEOBJ_TEXT else 'other',
                    'bounds':list(obj.get_bounds())}
            if item['type'] == 'text':
                mode=raw.FPDFTextObj_GetTextRenderMode(obj)
                item.update(text=obj.extract(), font_size=obj.get_font_size(),
                            font=obj.get_font().get_base_name(),render_mode=mode,
                            editable=mode in {0,1,2},
                            unsupported_reason=None if mode in {0,1,2} else 'Invisible or clipping text is complex and not editable.')
            result.append(item)
        return result

def _inspect(path):
    with pdfium.PdfDocument(str(path)) as doc:
        pages = []
        for n in range(len(doc)):
            with closing(doc[n]) as page:
                pages.append({'page':n+1,'width':page.get_width(),'height':page.get_height(),
                              'objects':_objects(page)})
        return {'pages':pages, 'limitations':['Only top-level text objects are editable; nested form content is not listed as editable.']}

def _result(paths=(), **details):
    return {'outputs':[{'path':str(p),'name':p.name,'mime':mimetypes.guess_type(p.name)[0] or 'application/octet-stream'} for p in paths], 'details':details}

from contextvars import ContextVar

_ALLOCATED = ContextVar('pdf_outputs', default=None)

OPTION_KEYS = {
    'inspect_pdf': {'page'}, 'edit_pdf_text': {'page','object_index','text','font','font_size'},
    'annotate_pdf': {'page','kind','text','rect'}, 'form_template': {'template_id'},
    'protect_pdf': {'password','owner_password'}, 'unlock_pdf': {'password'},
    'remove_watermark': {'exact_text','page','annotation_subtype'},
    'compress_pdf': {'mode','dpi','quality'}, 'convert_to_pdf': {'timeout'},
    'convert_from_pdf': {'format','dpi','quality'},
    'ocr_pdf': {'language','dpi','timeout','include_editable_text'},
    'ocr_image': {'language','dpi','timeout','include_editable_text'},
}


def run(request):
    allocated=[]
    token=_ALLOCATED.set(allocated)
    try:
        result=_run(request)
        directory=Path(request['output_dir']).resolve()
        for output in result['outputs']:
            path=Path(output['path'])
            if path.parent != directory or path.is_symlink() or not path.is_file() or not path.stat().st_size:
                raise PdfOperationError('Processing did not produce a valid file inside the job directory.')
        return result
    except Exception as exc:
        for path in allocated: path.unlink(missing_ok=True)
        if isinstance(exc,PdfOperationError): raise
        raise PdfOperationError('Document processing failed. Check the file format, encryption, and selected options.') from None
    finally:
        _ALLOCATED.reset(token)

def _run(request):
    operation = request.get('operation')
    if operation not in OPERATIONS:
        raise PdfOperationError('Unsupported native PDF operation.')
    inputs = _inputs(request)
    directory = Path(request.get('output_dir', ''))
    if not directory.is_absolute() or directory.is_symlink():
        raise PdfOperationError('output_dir must be an absolute job directory, not a symlink.')
    directory.mkdir(parents=True, exist_ok=True)
    directory = directory.resolve()
    if not inputs and operation != 'form_template':
        raise PdfOperationError('Select one input file.')
    if len(inputs) > 1:
        raise PdfOperationError('This operation accepts one input file per job.')
    options = request.get('options', {})
    if not isinstance(options, dict):
        raise PdfOperationError('options must be an object.')
    unknown=set(options)-OPTION_KEYS[operation]
    if unknown:
        # Do not echo caller values, which may contain document contents or passwords.
        raise PdfOperationError('Unknown option for this operation. Check PDF-OPTIONS.md.')
    if operation == 'inspect_pdf':
        info=_inspect(inputs[0])
        if 'page' in options:
            index=_page(info['pages'],options)
            info['pages']=[info['pages'][index]]
        return _result(**info)
    if operation == 'edit_pdf_text':
        return _edit(inputs[0], directory, options)
    if operation in {'protect_pdf', 'unlock_pdf'}:
        from pypdf import PdfReader, PdfWriter
        password = options.get('password')
        if not isinstance(password, str) or not password:
            raise PdfOperationError('A non-empty password is required.')
        reader = PdfReader(inputs[0])
        if reader.is_encrypted:
            if not reader.decrypt(password):
                raise PdfOperationError('Incorrect password; use a known document password.')
        elif operation == 'unlock_pdf':
            raise PdfOperationError('This PDF is not password encrypted.')
        writer = PdfWriter(clone_from=reader)
        if operation == 'protect_pdf':
            owner = options.get('owner_password')
            if not isinstance(owner, str) or not owner or owner == password:
                raise PdfOperationError('A distinct non-empty owner_password is required.')
            writer.encrypt(password, owner, algorithm='AES-256')
        dest = _target(directory, '.pdf')
        writer.write(dest)
        return _result([dest], encryption='AES-256' if operation=='protect_pdf' else 'none',
                       limitations=['Owner permissions are advisory. Keep passwords separately; no password recovery is provided.'])
    if operation == 'remove_watermark':
        return _remove_watermark(inputs[0], directory, options)
    if operation == 'form_template':
        from reportlab.pdfgen import canvas
        templates = {'contact':['name','email','message'], 'approval':['requester','description','approved'],
                     'invoice':['invoice_number','customer','amount']}
        template = options.get('template_id', 'contact')
        if template not in templates:
            raise PdfOperationError('template_id must be contact, approval, or invoice.')
        dest = _target(directory, '.pdf')
        c = canvas.Canvas(str(dest), pagesize=(612,792))
        c.setFont('Helvetica-Bold', 20)
        c.drawString(54,738,template.title() + ' form')
        c.setFont('Helvetica',10)
        c.drawString(54,710,'General-purpose form; not legal or accounting advice.')
        for i, name in enumerate(templates[template]):
            y = 640 - i * 100
            c.drawString(54,y+28,name.replace('_',' ').title())
            if name == 'approved':
                c.acroForm.checkbox(name=name,x=54,y=y,buttonStyle='check',size=18)
            else:
                c.acroForm.textfield(name=name,x=54,y=y,width=480,height=24,
                                     fontName='Helvetica',fontSize=11)
        c.showPage(); c.save()
        return _result([dest], template_id=template, limitations=['General-purpose fillable sample, not a legal template. Standard Helvetica form font; non-Latin field display depends on the viewer.'])
    if operation == 'annotate_pdf':
        import math
        from pypdf import PdfWriter
        from pypdf.annotations import Text, Highlight
        from pypdf.generic import ArrayObject, FloatObject, NameObject, TextStringObject
        writer = PdfWriter(clone_from=inputs[0])
        index = _page(writer.pages, options)
        rect = options.get('rect')
        if not isinstance(rect, list) or len(rect) != 4 or not all(isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v) for v in rect):
            raise PdfOperationError('rect must be four finite PDF coordinates [x1,y1,x2,y2].')
        x1,y1,x2,y2 = rect
        if x1 >= x2 or y1 >= y2:
            raise PdfOperationError('rect must have positive width and height.')
        text = options.get('text', '')
        if not isinstance(text,str) or len(text)>10000:
            raise PdfOperationError('Annotation text must be at most 10000 characters.')
        kind = options.get('kind', 'comment')
        if kind == 'comment': annotation = Text(rect=rect,text=text)
        elif kind == 'highlight':
            annotation = Highlight(rect=rect,quad_points=ArrayObject([FloatObject(v) for v in [x1,y2,x2,y2,x1,y1,x2,y1]]))
            annotation[NameObject('/Contents')] = TextStringObject(text)
        else: raise PdfOperationError('kind must be comment or highlight.')
        writer.add_annotation(index, annotation)
        dest = _target(directory, '.pdf'); writer.write(dest)
        return _result([dest], kind=kind, limitations=['Real non-flattened annotations. Coordinates are PDF user space; rotated/cropped pages require UI coordinate mapping. Viewer rendering varies.'])
    if operation == 'compress_pdf':
        from pypdf import PdfWriter
        mode = options.get('mode','lossless')
        dest = _target(directory, '.pdf')
        if mode == 'lossless':
            writer = PdfWriter(clone_from=inputs[0])
            for page in writer.pages: page.compress_content_streams()
            writer.compress_identical_objects(remove_duplicates=True,remove_unreferenced=True)
            writer.write(dest)
            limitations = ['Lossless stream compression/deduplication; already compressed files may grow. Text, annotations and forms are retained.']
        elif mode == 'raster':
            from reportlab.pdfgen import canvas
            from reportlab.lib.utils import ImageReader
            import io
            dpi = _integer(options.get('dpi',150),72,300,'dpi')
            quality = _integer(options.get('quality',75),20,95,'quality')
            c = canvas.Canvas(str(dest))
            with pdfium.PdfDocument(str(inputs[0])) as doc:
                for n in range(len(doc)):
                    with closing(doc[n]) as page:
                        image = _render(page,dpi)
                        buffer = io.BytesIO(); image.convert('RGB').save(buffer,format='JPEG',quality=quality)
                        w,h = page.get_size(); c.setPageSize((w,h))
                        c.drawImage(ImageReader(buffer),0,0,width=w,height=h)
                        c.showPage()
            c.save()
            limitations = ['Lossy raster compression flattens text, links, annotations and forms into pixels. Text is no longer searchable or editable. Size reduction is not guaranteed.']
        else: raise PdfOperationError('mode must be lossless or raster.')
        return _result([dest],mode=mode,input_bytes=inputs[0].stat().st_size,output_bytes=dest.stat().st_size,limitations=limitations)
    if operation == 'convert_from_pdf':
        return _export(inputs[0], directory, options)
    if operation == 'convert_to_pdf':
        return _to_pdf(inputs[0], directory, options, request.get('engines', {}))
    if operation in {'ocr_pdf','ocr_image'}:
        _engine(request.get('engines',{}),'tesseract','Tesseract')
        return _ocr(inputs[0],directory,options,request.get('engines',{}),operation)
    raise PdfOperationError('Unsupported native PDF operation.')


def _ocr(path, directory, options, engines, operation):
    import re
    from pypdf import PdfWriter, PdfReader
    from PIL import Image, ImageSequence, ImageOps
    engine=_engine(engines,'tesseract','Tesseract')
    language=options.get('language','eng')
    if not isinstance(language,str) or not re.fullmatch(r'[A-Za-z0-9_]+(?:\+[A-Za-z0-9_]+)*',language):
        raise PdfOperationError('language must be an installed Tesseract language code, such as eng or eng+deu.')
    dpi=_integer(options.get('dpi',300),72,300,'dpi')
    timeout=_integer(options.get('timeout',120),1,600,'timeout')
    include=options.get('include_editable_text',True)
    if not isinstance(include,bool): raise PdfOperationError('include_editable_text must be boolean.')
    if not include: raise PdfOperationError('OCR requires include_editable_text=true: an editable plain-text companion is always provided.')
    writer=PdfWriter(); texts=[]
    with tempfile.TemporaryDirectory(prefix='ocr-',dir=directory) as temporary:
        root=Path(temporary)
        def recognize(image,n):
            source=root/f'page-{n}.png'; image.convert('RGB').save(source)
            base=root/f'ocr-{n}'
            _subprocess([str(engine),str(source),str(base),'-l',language,'--dpi',str(dpi),'pdf','txt'],timeout,'Tesseract')
            pdf=base.with_suffix('.pdf'); txt=base.with_suffix('.txt')
            if not pdf.is_file() or not txt.is_file():
                raise PdfOperationError('Tesseract did not produce both searchable PDF and editable text. Check PDF renderer support and language data.')
            reader=PdfReader(pdf)
            if not reader.pages: raise PdfOperationError('Tesseract produced an empty PDF.')
            writer.append(reader)
            texts.append(txt.read_text(encoding='utf-8'))
        if operation=='ocr_pdf':
            with pdfium.PdfDocument(str(path)) as doc:
                for n in range(len(doc)):
                    with closing(doc[n]) as page: recognize(_render(page,dpi),n)
        else:
            with Image.open(path) as image:
                for n,frame in enumerate(ImageSequence.Iterator(image)):
                    if frame.width*frame.height>40_000_000: raise PdfOperationError('Image exceeds 40 megapixels.')
                    recognize(ImageOps.exif_transpose(frame),n)
        dest=_target(directory,'.pdf'); writer.write(dest)
        txt=_target(directory,'.txt'); txt.write_text('\n\f\n'.join(texts),encoding='utf-8')
    return _result([dest,txt],language=language,page_count=len(texts),editable_text_format='txt',
                   limitations=['OCR is probabilistic; review recognized text. Searchable PDF retains raster scans with a recognition layer, not semantically editable original layout. The separate UTF-8 plain-text output is editable. Existing annotations/forms are flattened by rendering; physical page size may change for image inputs.'])


def _safe_html_text(value):
    from html.parser import HTMLParser
    class TextOnly(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts=[]; self.ignored=0
        def handle_starttag(self,tag,attrs):
            if tag in {'script','style','iframe','object','svg','template'}: self.ignored+=1
            if tag in {'p','div','h1','h2','h3','br','li','tr'} and not self.ignored: self.parts.append('\n')
        def handle_endtag(self,tag):
            if tag in {'script','style','iframe','object','svg','template'}: self.ignored=max(0,self.ignored-1)
            if tag in {'p','div','h1','h2','h3','li','tr'} and not self.ignored: self.parts.append('\n')
        def handle_data(self,data):
            if not self.ignored: self.parts.append(data)
    parser=TextOnly(); parser.feed(value)
    return ''.join(parser.parts)


def _to_pdf(path, directory, options, engines):
    from reportlab.pdfgen import canvas
    from reportlab.platypus import SimpleDocTemplate, Preformatted
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.utils import ImageReader
    from PIL import Image, ImageOps, ImageSequence
    suffix=path.suffix.lower()
    dest=_target(directory,'.pdf')
    if suffix in {'.txt','.html','.htm'}:
        if path.stat().st_size > 5*1024*1024:
            raise PdfOperationError('Text/HTML input exceeds the 5 MiB limit.')
        text=path.read_text(encoding='utf-8-sig')
        if suffix != '.txt': text=_safe_html_text(text)
        # No HTML markup is passed to a browser or renderer; assets/URLs are never loaded.
        import textwrap
        text='\n'.join('\n'.join(textwrap.wrap(line,90,replace_whitespace=False) or ['']) for line in text.splitlines())
        style=getSampleStyleSheet()['Code']; style.fontName='Courier'; style.fontSize=10; style.leading=13
        doc=SimpleDocTemplate(str(dest))
        doc.build([Preformatted(text or ' ',style)])
        limitations=['Plain UTF-8 text laid out with standard Courier; complex Unicode glyph coverage is limited. HTML is text-only: scripts/styles/embedded content and all asset loading are discarded, not browser-rendered.']
    elif suffix in {'.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'}:
        c=canvas.Canvas(str(dest))
        with Image.open(path) as image:
            for frame in ImageSequence.Iterator(image):
                if frame.width*frame.height>40_000_000: raise PdfOperationError('Image exceeds 40 megapixels.')
                frame=ImageOps.exif_transpose(frame).convert('RGB')
                w,h=frame.size
                scale=min(540/w,720/h,1)
                c.setPageSize((w*scale,h*scale)); c.drawImage(ImageReader(frame),0,0,width=w*scale,height=h*scale)
                c.showPage()
        c.save()
        limitations=['Images become PDF pages; no searchable text is added. Transparent pixels are converted to RGB.']
    elif suffix in {'.docx','.xlsx','.pptx','.odt','.ods','.odp'}:
        _validate_office(path)
        engine = _engine(engines,'libreoffice','LibreOffice')
        timeout = _integer(options.get('timeout',120),1,600,'timeout')
        import shutil
        with tempfile.TemporaryDirectory(prefix='office-',dir=directory) as temporary:
            root=Path(temporary); profile=root/'profile'; (profile/'user').mkdir(parents=True)
            (profile/'user'/'registrymodifications.xcu').write_text(
                '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item><item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>1</value></prop></item></oor:items>')
            staged=root/('input'+suffix); shutil.copyfile(path,staged)
            _subprocess([str(engine),'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--nodefault','--nolockcheck','--norestore','--convert-to','pdf','--outdir',str(root),str(staged)],timeout,'LibreOffice')
            produced=root/'input.pdf'
            if not produced.is_file(): raise PdfOperationError('LibreOffice produced no PDF; verify the document and installed export filters.')
            shutil.copyfile(produced,dest)
        limitations=['LibreOffice layout/fonts may differ from Microsoft Office. Macros, embedded executables and external relationships are rejected. Engine must be explicitly installed with consent. A hardened OS sandbox remains recommended for untrusted Office documents.']
    else:
        raise PdfOperationError('Unsupported input. Choose UTF-8 text, HTML, a supported image, or DOCX/XLSX/PPTX/ODT/ODS/ODP.')
    from pypdf import PdfReader
    if not PdfReader(dest).pages: raise PdfOperationError('PDF conversion produced no pages.')
    return _result([dest],limitations=limitations)


def _engine(engines, key, label):
    value = engines.get(key) if isinstance(engines,dict) else None
    if not isinstance(value,str) or not value:
        raise PdfOperationError(f'Install the optional {label} component with explicit consent before using this operation.')
    path=Path(value)
    if not path.is_absolute() or not path.is_file() or not os.access(path,os.X_OK):
        raise PdfOperationError(f'The installed {label} executable is unavailable. Reinstall the component.')
    return path


def _subprocess(args, timeout, label):
    import subprocess
    try:
        result=subprocess.run(args,stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL,timeout=timeout,shell=False,check=False)
    except subprocess.TimeoutExpired:
        raise PdfOperationError(f'{label} timed out. Try a smaller document or increase timeout.') from None
    except OSError:
        raise PdfOperationError(f'{label} could not start. Check the installed component and its runtime libraries.') from None
    if result.returncode:
        raise PdfOperationError(f'{label} failed. Check document validity, installed language data/export filters, and component compatibility.')


def _validate_office(path):
    import zipfile
    import stat
    import xml.etree.ElementTree as ET
    from pathlib import PurePosixPath
    try:
        with zipfile.ZipFile(path) as archive:
            entries=archive.infolist()
            if len(entries)>10000 or sum(e.file_size for e in entries)>200*1024*1024:
                raise PdfOperationError('Unsafe Office archive: expanded size or entry count exceeds limits.')
            for entry in entries:
                name=entry.filename
                lower=name.lower()
                if name.startswith('/') or '\\' in name or '..' in PurePosixPath(name).parts or ':' in name or stat.S_ISLNK(entry.external_attr>>16):
                    raise PdfOperationError('Unsafe Office archive: traversal or symbolic-link entry.')
                if entry.flag_bits & 1 or entry.file_size > 30*1024*1024 or entry.file_size > max(1,entry.compress_size)*500:
                    raise PdfOperationError('Unsafe Office archive: encrypted, oversized or suspicious compressed entry.')
                if any(token in lower for token in ['vbaproject','/embeddings/','scripts/','basic/','oleobject']):
                    raise PdfOperationError('Unsafe Office archive: macros and embedded executable objects are not supported.')
                if lower.endswith(('.xml','.rels')):
                    value=archive.read(entry)
                    if b'<!DOCTYPE' in value.upper() or b'<!ENTITY' in value.upper():
                        raise PdfOperationError('Unsafe Office archive: XML entity declarations are not allowed.')
                    root=ET.fromstring(value)
                    for node in root.iter():
                        if node.attrib.get('TargetMode','').lower()=='external':
                            raise PdfOperationError('Unsafe Office archive: external relationships are not allowed.')
                        for key,target in node.attrib.items():
                            if key.endswith('}href') and (':' in target or target.startswith(('/','\\','..'))):
                                raise PdfOperationError('Unsafe Office archive: external resource references are not allowed.')
                        if node.tag.rsplit('}',1)[-1].lower() in {'script','event-listener','dde-link','dde-connection','oleobject'}:
                            raise PdfOperationError('Unsafe Office archive: active content is not supported.')
                        if node.text and any(token in node.text.upper() for token in ['DDEAUTO','INCLUDETEXT','INCLUDEPICTURE']):
                            raise PdfOperationError('Unsafe Office archive: active document fields are not supported.')
    except (zipfile.BadZipFile,ET.ParseError):
        raise PdfOperationError('Unsafe Office archive: invalid ZIP or XML.') from None


def _extracted(path):
    import pdfplumber
    with pdfplumber.open(path) as doc:
        return [{'text':page.extract_text(layout=False) or '',
                 'words':page.extract_words(), 'tables':page.extract_tables(),
                 'width':page.width, 'height':page.height} for page in doc.pages]


def _spreadsheet_value(value):
    # Avoid spreadsheet formula injection from untrusted document text.
    value = '' if value is None else str(value)
    return "'" + value if value.startswith(('=','+','-','@','\t','\r')) else value


def _export(path, directory, options):
    import csv
    import html
    format = options.get('format','docx')
    if format not in {'docx','xlsx','pptx','png','jpeg','txt','csv','html'}:
        raise PdfOperationError('format must be docx, xlsx, pptx, png, jpeg, txt, csv, or html.')
    if format in {'png','jpeg'}:
        dpi = _integer(options.get('dpi',150),72,300,'dpi')
        quality = _integer(options.get('quality',85),20,95,'quality')
        paths=[]
        with pdfium.PdfDocument(str(path)) as doc:
            for n in range(len(doc)):
                with closing(doc[n]) as page:
                    image=_render(page,dpi)
                    dest=_target(directory,'.'+format)
                    image.convert('RGB').save(dest,format='JPEG' if format=='jpeg' else 'PNG',**({'quality':quality} if format=='jpeg' else {}))
                    paths.append(dest)
        return _result(paths,format=format,limitations=['Raster image export; no editable text. One output per PDF page, in page order.'])
    pages = _extracted(path)
    dest = _target(directory,'.'+format)
    if format == 'txt':
        dest.write_text('\n\f\n'.join(p['text'] for p in pages),encoding='utf-8')
    elif format == 'html':
        body=''.join('<section><h2>Page '+str(n+1)+'</h2><pre>'+html.escape(p['text'])+'</pre></section>' for n,p in enumerate(pages))
        dest.write_text('<!doctype html><html><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'"></head><body>'+body+'</body></html>',encoding='utf-8')
    elif format == 'csv':
        with dest.open('w',encoding='utf-8',newline='') as f:
            writer=csv.writer(f)
            for p in pages:
                rows=[row for table in p['tables'] for row in table] if p['tables'] else [[line] for line in p['text'].splitlines()]
                writer.writerows([[_spreadsheet_value(v) for v in row] for row in rows])
    elif format == 'docx':
        from docx import Document
        from docx.shared import Pt
        doc=Document()
        doc.styles['Normal'].font.size=Pt(11)
        for n,p in enumerate(pages):
            if n: doc.add_page_break()
            for line in p['text'].splitlines(): doc.add_paragraph(line)
            for table in p['tables']:
                if not table: continue
                t=doc.add_table(rows=0,cols=max(map(len,table)))
                t.style='Table Grid'
                for row in table:
                    cells=t.add_row().cells
                    for i,value in enumerate(row): cells[i].text=value or ''
        doc.save(dest)
    elif format == 'xlsx':
        from openpyxl import Workbook
        doc=Workbook(); doc.remove(doc.active)
        for n,p in enumerate(pages):
            sheet=doc.create_sheet('Page '+str(n+1))
            rows=[row for table in p['tables'] for row in table] if p['tables'] else [[line] for line in p['text'].splitlines()]
            for row in rows: sheet.append([_spreadsheet_value(v) for v in row])
            sheet.column_dimensions['A'].width=60
        doc.save(dest)
    elif format == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        doc=Presentation()
        doc.slide_width=Inches(10)
        doc.slide_height=Inches(7.5)
        for p in pages:
            slide=doc.slides.add_slide(doc.slide_layouts[6])
            # Group same-baseline words into editable lines, retaining approximate positions.
            lines=[]
            for word in p['words']:
                if lines and abs(lines[-1]['top']-word['top'])<3:
                    lines[-1]['text'] += ' '+word['text']
                    lines[-1]['x1'] = max(lines[-1]['x1'],word['x1'])
                    lines[-1]['bottom'] = max(lines[-1]['bottom'],word['bottom'])
                else: lines.append(dict(word))
            for word in lines:
                left=int(word['x0']/p['width']*doc.slide_width)
                top=int(word['top']/p['height']*doc.slide_height)
                width=max(int((word['x1']-word['x0'])/p['width']*doc.slide_width)+Inches(.1),Inches(.2))
                height=max(int((word['bottom']-word['top'])/p['height']*doc.slide_height)+Inches(.1),Inches(.2))
                box=slide.shapes.add_textbox(left,top,width,height)
                box.text_frame.margin_left=box.text_frame.margin_right=0
                box.text=word['text']
                box.text_frame.paragraphs[0].font.size=Pt(max(6,min(48,(word['bottom']-word['top'])*.8)))
            # Full text in speaker notes makes reading order accessible without raster flattening.
            slide.notes_slide.notes_text_frame.text=p['text']
        doc.save(dest)
    return _result([dest],format=format,page_count=len(pages),limitations=['Best-effort reconstruction, not an Office round-trip. Tables use line detection; merged/borderless cells may be wrong. DOCX includes extracted lines plus detected tables; table text may repeat. PPTX uses positioned editable words, not original typography/images. Scans need OCR first.'])


def _render(page, dpi):
    width,height=page.get_size()
    if width * height * (dpi/72)**2 > 40_000_000:
        raise PdfOperationError('Rendered page exceeds 40 megapixels; reduce dpi.')
    with closing(page.render(scale=dpi/72)) as bitmap:
        return bitmap.to_pil().copy()


def _remove_watermark(path, directory, options):
    exact = options.get('exact_text')
    subtype = options.get('annotation_subtype')
    if exact is not None and (not isinstance(exact, str) or not exact):
        raise PdfOperationError('exact_text must be a non-empty exact object match.')
    if subtype is not None and subtype not in {'Watermark','Stamp','FreeText','Text','Highlight'}:
        raise PdfOperationError('Select a supported annotation_subtype explicitly.')
    if exact is None and subtype is None:
        raise PdfOperationError('Select exact_text and/or annotation_subtype; automatic watermark guessing is not supported.')
    removed = 0
    with pdfium.PdfDocument(str(path)) as doc:
        indices = [_page(doc, options)] if 'page' in options else range(len(doc))
        for n in indices:
            with closing(doc[n]) as page:
                if exact is not None:
                    with closing(page.get_textpage()) as tp:
                        matches = [obj for obj in page.get_objects(max_depth=1,textpage=tp)
                                   if obj.type == raw.FPDF_PAGEOBJ_TEXT and obj.extract() == exact]
                    for obj in matches:
                        if not raw.FPDFPage_RemoveObject(page, obj):
                            raise PdfOperationError('PDFium could not remove a matching object.')
                        raw.FPDFPageObj_Destroy(obj)
                        removed += 1
                    if matches: page.gen_content()
        dest = _target(directory, '.pdf')
        doc.save(str(dest))
    if subtype:
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import ArrayObject, NameObject
        writer = PdfWriter(clone_from=dest)
        for n in indices:
            page = writer.pages[n]
            kept = []
            for reference in page.get('/Annots', []):
                annotation = reference.get_object()
                if annotation.get('/Subtype') == '/' + subtype:
                    removed += 1
                else: kept.append(reference)
            if '/Annots' in page: page[NameObject('/Annots')] = ArrayObject(kept)
        writer.write(dest)
    if not removed:
        dest.unlink(missing_ok=True)
        raise PdfOperationError('No exact watermark match found. Flattened images and nested form text are not automatically removable.')
    return _result([dest], removed_objects=removed, limitations=['Only exact top-level text objects or explicitly selected annotation subtypes are removed. Other content is not guessed. This is not secure redaction.'])


def _target(directory, suffix):
    # Unique names, exclusive reservation: never overwrite inputs or past results.
    fd, name = tempfile.mkstemp(prefix='paperwork-', suffix=suffix, dir=directory)
    os.close(fd)
    path=Path(name)
    allocated=_ALLOCATED.get()
    if allocated is not None: allocated.append(path)
    return path


def _integer(value, minimum, maximum, name):
    if isinstance(value, bool) or not isinstance(value, int) or not minimum <= value <= maximum:
        raise PdfOperationError(f'{name} must be an integer from {minimum} to {maximum}.')
    return value


def _page(doc, options):
    return _integer(options.get('page', 1), 1, len(doc), 'page') - 1


def _edit(path, directory, options):
    import ctypes
    text = options.get('text')
    if not isinstance(text, str) or '\x00' in text or len(text) > 10000:
        raise PdfOperationError('text must contain at most 10000 characters and no NUL.')
    if any(ord(c)>255 for c in text):
        raise PdfOperationError('Complex-script and non-Latin-1 text replacement is unsupported; use a full document editor for shaping/font embedding.')
    with pdfium.PdfDocument(str(path)) as doc:
        with closing(doc[_page(doc, options)]) as page:
            objects = list(page.get_objects(max_depth=1))
            index = _integer(options.get('object_index'), 0, len(objects)-1, 'object_index')
            obj = objects[index]
            if obj.type != raw.FPDF_PAGEOBJ_TEXT:
                raise PdfOperationError('Select a top-level text object; nested forms and complex objects are unsupported.')
            if raw.FPDFTextObj_GetTextRenderMode(obj) not in {0,1,2}:
                raise PdfOperationError('Invisible or clipping text is complex and unsupported for editing.')
            size = options.get('font_size', obj.get_font_size())
            if isinstance(size, bool) or not isinstance(size, (float, int)) or not 1 <= size <= 500:
                raise PdfOperationError('font_size must be between 1 and 500 points.')
            font = options.get('font')
            if font is not None:
                if font not in {'Helvetica','Helvetica-Bold','Helvetica-Oblique','Times-Roman','Times-Bold','Times-Italic','Courier','Courier-Bold','Courier-Oblique'}:
                    raise PdfOperationError('Choose a supported standard PDF font.')
                if any(ord(c)>255 for c in text):
                    raise PdfOperationError('Standard-font substitution supports Latin-1 text only.')
                # PDFium has no font setter: replace the actual text object, preserving matrix and colors.
                replacement = raw.FPDFPageObj_NewTextObj(doc, font.encode('ascii'), size)
                if not replacement:
                    raise PdfOperationError('PDFium could not create the selected font.')
                matrix = raw.FS_MATRIX()
                raw.FPDFPageObj_GetMatrix(obj, matrix)
                raw.FPDFPageObj_SetMatrix(replacement, matrix)
                for get, set_ in [(raw.FPDFPageObj_GetFillColor, raw.FPDFPageObj_SetFillColor),
                                  (raw.FPDFPageObj_GetStrokeColor, raw.FPDFPageObj_SetStrokeColor)]:
                    rgba = [ctypes.c_uint() for _ in range(4)]
                    if get(obj, *[ctypes.byref(v) for v in rgba]):
                        set_(replacement, *[v.value for v in rgba])
                raw.FPDFTextObj_SetTextRenderMode(replacement, raw.FPDFTextObj_GetTextRenderMode(obj))
                target = replacement
            else:
                target = obj
                raw.FPDFTextObj_SetFontSize(target, size)
            encoded = text.encode('utf-16-le') + b'\0\0'
            buffer = (ctypes.c_ushort * (len(encoded)//2)).from_buffer_copy(encoded)
            if not raw.FPDFText_SetText(target, buffer):
                if font: raw.FPDFPageObj_Destroy(target)
                raise PdfOperationError('The selected font cannot encode this text; choose a standard font for simple Latin text.')
            if font:
                if not raw.FPDFPage_RemoveObject(page, obj):
                    raw.FPDFPageObj_Destroy(target)
                    raise PdfOperationError('PDFium could not remove the original object.')
                raw.FPDFPageObj_Destroy(obj)
                raw.FPDFPage_InsertObject(page, target)
            page.gen_content()
        dest = _target(directory, '.pdf')
        doc.save(str(dest))
    return _result([dest], limitations=['No reflow. Font substitution replaces the text object and preserves its matrix/colors, not clipping, blend state, kerning or drawing order. Complex scripts and subset-font encoding may require a standard Latin font.'])
