"""Exercise python-pptx resource lookup with the freezer's actual data layout."""
from pathlib import Path
import runpy
import shutil
import zipfile

import pytest

ROOT = Path(__file__).resolve().parents[2]


def test_frozen_pptx_data_layout_supports_speaker_notes(tmp_path, monkeypatch):
    # Run with the packaging venv. No PyInstaller executable build is required
    # for this tight regression; the real frozen-worker acceptance is separate.
    hooks = pytest.importorskip("PyInstaller.utils.hooks")
    pptx = pytest.importorskip("pptx")
    import pptx.api
    import pptx.oxml

    command = runpy.run_path(str(ROOT / "scripts/freeze-worker.py"))["build_command"]()
    datas = []
    for option, package in zip(command, command[1:]):
        if package != "pptx":
            continue
        if option == "--collect-data":
            datas.extend(hooks.collect_data_files(package))
        elif option == "--collect-all":
            datas.extend(hooks.collect_all(package)[0])
    assert datas, "Freezer must collect the PPTX templates"
    internal = tmp_path / "_internal"
    for source, destination in datas:
        target = internal / destination
        if Path(source).is_dir():
            shutil.copytree(source, target, dirs_exist_ok=True)
        else:
            target.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source, target / Path(source).name)

    # PyInstaller assigns these paths even when the modules live only in PYZ.
    # Exercise the dependency's real unnormalized oxml/../templates lookup,
    # not merely the existence of templates/notesMaster.xml.
    monkeypatch.setattr(pptx.api, "__file__", str(internal / "pptx/api.py"))
    monkeypatch.setattr(pptx.oxml, "__file__", str(internal / "pptx/oxml/__init__.py"))
    document = pptx.Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "Original invoice"
    slide.notes_slide.notes_text_frame.text = "Original invoice\nDRAFT"
    output = tmp_path / "synthetic.pptx"
    document.save(output)

    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None
        assert b"Original invoice" in archive.read("ppt/slides/slide1.xml")
        assert b"DRAFT" in archive.read("ppt/notesSlides/notesSlide1.xml")
        assert "ppt/notesMasters/notesMaster1.xml" in archive.namelist()
    reopened = pptx.Presentation(output)
    assert reopened.slides[0].notes_slide.notes_text_frame.text == "Original invoice\nDRAFT"
